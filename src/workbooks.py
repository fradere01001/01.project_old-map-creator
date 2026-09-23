"""Inspection and safety metadata for existing XLSX destinations."""

from __future__ import annotations

import hashlib
import math
from pathlib import Path

from openpyxl import load_workbook

from .models import (
    DestinationMode,
    DestinationPlan,
    WorkbookLayout,
    location_identity,
    normalize_header,
)


class WorkbookInspectionError(ValueError):
    pass


LOCATION_HEADERS = {"address", "latitude", "longitude"}
REPORT_HEADERS = (
    "Source Row", "Input Address", "Status", "Cache Hit", "Duplicate Of",
    "Resolved Address", "Latitude", "Longitude", "Error",
)


def file_fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def workbook_location_identity(
    address: object, latitude: object, longitude: object
) -> str | None:
    """Return an identity only for a usable workbook location row."""
    try:
        if isinstance(latitude, bool) or isinstance(longitude, bool):
            return None
        numeric_latitude = float(latitude)
        numeric_longitude = float(longitude)
    except (TypeError, ValueError):
        return None
    if not (math.isfinite(numeric_latitude) and math.isfinite(numeric_longitude)):
        return None
    if not (-90 <= numeric_latitude <= 90 and -180 <= numeric_longitude <= 180):
        return None
    return location_identity(
        None if address is None else str(address),
        numeric_latitude,
        numeric_longitude,
    )


def _new_sheet_name(existing: list[str], preferred: str, fallback: str) -> str:
    if preferred not in existing:
        return preferred
    if fallback not in existing:
        return fallback
    number = 2
    while f"{fallback} {number}" in existing:
        number += 1
    return f"{fallback} {number}"


def _header_mapping(sheet: object) -> tuple[int, dict[str, int]] | None:
    for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = [normalize_header(value) for value in cells]
        if not any(values):
            continue
        mapping: dict[str, int] = {}
        for required in LOCATION_HEADERS:
            positions = [index + 1 for index, value in enumerate(values) if value == required]
            if len(positions) != 1:
                return None
            mapping[required] = positions[0]
        return row_number, mapping
    return None


def inspect_existing_workbook(path: str | Path) -> DestinationPlan:
    candidate = Path(path).expanduser().resolve()
    if candidate.suffix.lower() != ".xlsx":
        raise WorkbookInspectionError("Only existing .xlsx workbooks can be extended.")
    if not candidate.is_file():
        raise WorkbookInspectionError(f"Existing workbook is not readable: {candidate}")
    try:
        before = file_fingerprint(candidate)
    except OSError as error:
        raise WorkbookInspectionError(
            f"Existing workbook is not readable: {candidate}: {error}"
        ) from error
    try:
        workbook = load_workbook(candidate, data_only=False)
    except Exception as error:
        raise WorkbookInspectionError(f"Cannot safely load existing workbook: {error}") from error
    try:
        matches: list[tuple[str, int, dict[str, int]]] = []
        for sheet in workbook.worksheets:
            mapped = _header_mapping(sheet)
            if mapped:
                matches.append((sheet.title, mapped[0], mapped[1]))
        named = [item for item in matches if item[0] == "Locations"]
        selected = named[0] if len(named) == 1 else matches[0] if len(matches) == 1 else None
        if selected:
            sheet_name, header_row, columns = selected
            recognized = True
        else:
            sheet_name = _new_sheet_name(workbook.sheetnames, "Locations", "Old Map Locations")
            header_row = 1
            columns = {"address": 1, "latitude": 2, "longitude": 3}
            recognized = False

        identities: list[str] = []
        if recognized:
            sheet = workbook[sheet_name]
            for row in range(header_row + 1, sheet.max_row + 1):
                identity = workbook_location_identity(
                    sheet.cell(row, columns["address"]).value,
                    sheet.cell(row, columns["latitude"]).value,
                    sheet.cell(row, columns["longitude"]).value,
                )
                if identity:
                    identities.append(identity)

        report_name = "Import Report"
        report_start = 2
        if report_name in workbook.sheetnames:
            report = workbook[report_name]
            headers = tuple(report.cell(1, index).value for index in range(1, 10))
            if headers == REPORT_HEADERS:
                report_start = report.max_row + 1
            else:
                report_name = _new_sheet_name(workbook.sheetnames, report_name, "Old Map Import Report")
        layout = WorkbookLayout(
            sheet_name, header_row, columns["address"], columns["latitude"],
            columns["longitude"], recognized, report_name, report_start,
        )
        return DestinationPlan(
            DestinationMode.EXTEND, candidate, "xlsx", layout, before,
            tuple(dict.fromkeys(identities)),
        )
    finally:
        workbook.close()


def extension_preview(plan: DestinationPlan) -> str:
    assert plan.layout is not None
    action = "append to recognized" if plan.layout.recognized else "create"
    return (
        f"Existing workbook: {plan.path}\n"
        f"Layout: {action} worksheet '{plan.layout.location_sheet}'\n"
        f"Existing unique locations: {len(plan.existing_location_identities)}\n"
        "Existing worksheets and unrelated cells will be preserved."
    )
