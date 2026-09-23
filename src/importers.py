"""CSV, text, and XLSX address source adapters."""

from __future__ import annotations

import csv
import hashlib
import io
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import ImportedData, SourceRow, normalize_header


class ImportSourceError(ValueError):
    pass


RECOGNIZED_ADDRESS_HEADERS = {"address", "location", "place", "indirizzo"}
SUPPORTED_INPUTS = {".csv", ".txt", ".xlsx"}


def source_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _require_source(path: Path) -> Path:
    resolved = path.expanduser().resolve()
    if resolved.suffix.lower() not in SUPPORTED_INPUTS:
        raise ImportSourceError(
            "Unsupported input format. Use CSV, TXT, or XLSX."
        )
    if not resolved.is_file():
        raise ImportSourceError(f"Input file is not readable: {resolved}")
    if resolved.stat().st_size == 0:
        raise ImportSourceError("Input file is empty.")
    return resolved


def _load_csv(path: Path) -> ImportedData:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeError as error:
        raise ImportSourceError("CSV input must be UTF-8 encoded.") from error
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text), dialect=dialect, strict=True))
    except csv.Error as error:
        raise ImportSourceError(f"Malformed CSV input: {error}") from error
    if not rows:
        raise ImportSourceError("CSV input is empty.")
    headers = tuple(str(value).strip() for value in rows[0])
    if not any(headers):
        raise ImportSourceError("CSV input has no header row.")
    source_rows = []
    for row_number, row in enumerate(rows[1:], start=2):
        padded = list(row) + [""] * max(0, len(headers) - len(row))
        values = {header: padded[index] for index, header in enumerate(headers)}
        source_rows.append(SourceRow(row_number, values))
    return ImportedData(path, source_hash(path), headers, tuple(source_rows))


def _load_txt(path: Path) -> ImportedData:
    try:
        lines = path.read_text(encoding="utf-8-sig").splitlines()
    except UnicodeError as error:
        raise ImportSourceError("TXT input must be UTF-8 encoded.") from error
    if not lines:
        raise ImportSourceError("TXT input is empty.")
    rows = tuple(
        SourceRow(number, {"address": line})
        for number, line in enumerate(lines, start=1)
    )
    return ImportedData(
        path,
        source_hash(path),
        ("address",),
        rows,
        text_mode=True,
    )


def list_xlsx_worksheets(path: str | Path) -> tuple[str, ...]:
    source = _require_source(Path(path))
    if source.suffix.lower() != ".xlsx":
        raise ImportSourceError("Worksheet selection is available only for XLSX.")
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:
        raise ImportSourceError(f"Cannot read XLSX input: {error}") from error
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


def _load_xlsx(path: Path, worksheet: str | None) -> ImportedData:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise ImportSourceError(f"Cannot read XLSX input: {error}") from error
    try:
        sheets = tuple(workbook.sheetnames)
        if not sheets:
            raise ImportSourceError("XLSX input has no worksheets.")
        if worksheet is None:
            if len(sheets) != 1:
                raise ImportSourceError(
                    "XLSX input has multiple worksheets; select one of: "
                    + ", ".join(sheets)
                )
            worksheet = sheets[0]
        if worksheet not in sheets:
            raise ImportSourceError(
                f"Worksheet '{worksheet}' not found. Available: {', '.join(sheets)}"
            )
        sheet = workbook[worksheet]
        header_row_number: int | None = None
        headers: tuple[str, ...] = ()
        data: list[SourceRow] = []
        for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = tuple("" if value is None else value for value in cells)
            if header_row_number is None:
                if not any(str(value).strip() for value in values):
                    continue
                header_row_number = row_number
                headers = tuple(str(value).strip() for value in values)
                if not any(headers):
                    raise ImportSourceError("XLSX worksheet has no header row.")
                continue
            padded = values + ("",) * max(0, len(headers) - len(values))
            mapped: dict[str, Any] = {
                header: padded[index] for index, header in enumerate(headers)
            }
            data.append(SourceRow(row_number, mapped))
        if header_row_number is None:
            raise ImportSourceError("XLSX worksheet is empty.")
        return ImportedData(
            path,
            source_hash(path),
            headers,
            tuple(data),
            worksheets=sheets,
            worksheet=worksheet,
        )
    finally:
        workbook.close()


def load_source(path: str | Path, worksheet: str | None = None) -> ImportedData:
    source = _require_source(Path(path))
    if source.suffix.lower() == ".csv":
        return _load_csv(source)
    if source.suffix.lower() == ".txt":
        return _load_txt(source)
    return _load_xlsx(source, worksheet)


def suggest_address_column(headers: tuple[str, ...]) -> str | None:
    matches = [
        header
        for header in headers
        if normalize_header(header) in RECOGNIZED_ADDRESS_HEADERS
    ]
    return matches[0] if len(matches) == 1 else None


def resolve_address_column(data: ImportedData, requested: str | None) -> str:
    if data.text_mode:
        return "address"
    if requested is not None:
        exact = [header for header in data.headers if header == requested]
        normalized = [
            header
            for header in data.headers
            if normalize_header(header) == normalize_header(requested)
        ]
        candidates = exact or normalized
        if len(candidates) == 1:
            return candidates[0]
        raise ImportSourceError(
            f"Address column '{requested}' not found. Available: "
            + ", ".join(data.headers)
        )
    suggestion = suggest_address_column(data.headers)
    if suggestion is None:
        raise ImportSourceError(
            "Address column is ambiguous. Available: " + ", ".join(data.headers)
        )
    return suggestion
